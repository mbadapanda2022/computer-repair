"""
Temporary Order Note tests — session-only Excel-like sheet with
bidirectional tax calculation and .xlsx export.

    python manage.py test accounting.tests_order_note -v 2
"""
from __future__ import annotations

from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import Client, TestCase, override_settings
from django.urls import reverse

User = get_user_model()


def post_rows(rows, supplier='Luna Mahanta', note=''):
    data = {'supplier': supplier, 'note': note}
    for key in ('date', 'product', 'qty', 'tax', 'price_excl', 'price_incl', 'price_src'):
        data[f'{key}[]'] = [row.get(key, '') for row in rows]
    return data


@override_settings(
    ALLOWED_HOSTS=['testserver', 'localhost', '127.0.0.1'],
    SECURE_SSL_REDIRECT=False,
)
class OrderNoteTests(TestCase):

    @classmethod
    def setUpTestData(cls):
        cls.staff = User.objects.create_user(
            username='note_staff', password='TestPass123!',
            is_staff=True, is_active=True,
        )

    def setUp(self):
        self.client = Client()
        self.client.force_login(self.staff)

    # ---- access ----
    def test_anonymous_redirected(self):
        anon = Client()
        r = anon.get('/order-note/')
        self.assertEqual(r.status_code, 302)
        self.assertIn('/login/', r['Location'].replace('%2F', '/'))

    def test_sheet_page_loads(self):
        r = self.client.get(reverse('accounting:order_note'))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'Order Note')
        self.assertContains(r, 'price_incl[]')

    # ---- saving (no validation, junk tolerated) ----
    def test_save_requires_nothing_and_drops_empty_rows(self):
        rows = [
            {'product': 'Mouse', 'qty': '2', 'price_excl': '250'},
            {'product': 'X', 'qty': 'junk', 'price_excl': ''},  # junk qty = 0
            {'date': '2026-10-17', 'product': '', 'qty': '', 'tax': '18',
             'price_excl': '', 'price_incl': '', 'price_src': 'excl'},  # prefilled-only row = empty
            {'product': '', 'qty': '0', 'price_excl': '', 'price_incl': ''},  # zero qty only = empty
        ]
        r = self.client.post(reverse('accounting:order_note'), post_rows(rows))
        self.assertEqual(r.status_code, 302)
        sheet = self.client.session['order_note']
        self.assertEqual(len(sheet['rows']), 2)
        self.assertEqual(sheet['rows'][0]['product'], 'Mouse')
        self.assertEqual(sheet['rows'][0]['price_incl'], '295.00')

    def test_excl_price_derives_incl_and_amount(self):
        rows = [{'product': 'SSD', 'qty': '3', 'tax': '18', 'price_excl': '1000'}]
        self.client.post(reverse('accounting:order_note'), post_rows(rows))
        row = self.client.session['order_note']['rows'][0]
        self.assertEqual(row['price_incl'], '1180.00')
        self.assertEqual(row['amount'], '3540.00')

    def test_incl_price_derives_excl(self):
        rows = [{'product': 'RAM', 'qty': '2', 'tax': '18',
                 'price_incl': '1180', 'price_src': 'incl'}]
        self.client.post(reverse('accounting:order_note'), post_rows(rows))
        row = self.client.session['order_note']['rows'][0]
        self.assertEqual(row['price_excl'], '1000.00')
        self.assertEqual(row['amount'], '2360.00')

    def test_default_tax_18_when_blank(self):
        rows = [{'product': 'Cable', 'qty': '1', 'tax': '', 'price_excl': '100'}]
        self.client.post(reverse('accounting:order_note'), post_rows(rows))
        row = self.client.session['order_note']['rows'][0]
        self.assertEqual(row['price_incl'], '118.00')

    def test_edit_and_delete_rows(self):
        self.client.post(reverse('accounting:order_note'), post_rows([
            {'product': 'A', 'qty': '1', 'price_excl': '100'},
            {'product': 'B', 'qty': '1', 'price_excl': '200'},
        ]))
        self.client.post(reverse('accounting:order_note'), post_rows([
            {'product': 'A edited', 'qty': '5', 'price_excl': '100'},
        ]))
        sheet = self.client.session['order_note']
        self.assertEqual(len(sheet['rows']), 1)
        self.assertEqual(sheet['rows'][0]['product'], 'A edited')

    def test_nothing_touches_database_tables(self):
        """Sheet lives only in the session — hard requirement from user."""
        from accounting.models import Purchase, Invoice
        before_p, before_i = Purchase.objects.count(), Invoice.objects.count()
        self.client.post(reverse('accounting:order_note'), post_rows([
            {'product': 'Monitor', 'qty': '1', 'price_excl': '9000'},
        ]))
        self.assertEqual(Purchase.objects.count(), before_p)
        self.assertEqual(Invoice.objects.count(), before_i)

    def test_clear_removes_sheet(self):
        self.client.post(reverse('accounting:order_note'), post_rows([
            {'product': 'A', 'qty': '1', 'price_excl': '100'},
        ]))
        r = self.client.post(reverse('accounting:order_note_clear'))
        self.assertEqual(r.status_code, 302)
        self.assertNotIn('order_note', self.client.session)

    # ---- excel export ----
    def test_excel_export_returns_workbook_with_totals(self):
        rows = [
            {'product': 'Mouse', 'qty': '2', 'price_excl': '250'},
            {'product': 'Keyboard', 'qty': '1', 'price_incl': '885', 'price_src': 'incl'},
        ]
        r = self.client.post(reverse('accounting:order_note_excel'), post_rows(rows))
        self.assertEqual(r.status_code, 200)
        self.assertIn('spreadsheetml', r['Content-Type'])
        self.assertIn('order-note_luna-mahanta_', r['Content-Disposition'])
        # saved as a side effect
        self.assertEqual(len(self.client.session['order_note']['rows']), 2)

        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(r.content))
        ws = wb.active
        self.assertIn('Order Note', ws['A1'].value)
        self.assertIn('Luna Mahanta', ws['A2'].value)
        # header row present
        self.assertEqual(ws.cell(row=5, column=3).value, 'Product / Item')
        # supplier-fill columns exist and are empty
        self.assertIn('Available', ws.cell(row=5, column=9).value)
        self.assertIsNone(ws.cell(row=6, column=9).value)
        # LIVE formulas so Excel recalculates when the supplier edits rates
        self.assertEqual(ws['F6'].value, 250)                       # typed excl stays a value
        self.assertEqual(ws['G6'].value, '=ROUND(F6*(1+E6/100),2)') # incl derived by formula
        self.assertEqual(ws['G7'].value, 885)                       # typed incl stays a value
        self.assertEqual(ws['F7'].value, '=ROUND(G7/(1+E7/100),2)') # excl derived by formula
        self.assertEqual(ws['H6'].value, '=ROUND(D6*G6,2)')          # amount by formula
        self.assertEqual(ws['H7'].value, '=ROUND(D7*G7,2)')
        self.assertEqual(ws['F8'].value, '=ROUND(SUMPRODUCT(D6:D7,F6:F7),2)')
        self.assertEqual(ws['F9'].value, '=ROUND(H10-F8,2)')
        self.assertEqual(ws['H10'].value, '=ROUND(SUM(H6:H7),2)')    # grand total
        # every table cell carries a border
        self.assertEqual(ws['I6'].border.left.style, 'medium')

    def test_excel_get_after_save_works(self):
        self.client.post(reverse('accounting:order_note'), post_rows([
            {'product': 'Pen', 'qty': '10', 'price_excl': '10'},
        ]))
        r = self.client.get(reverse('accounting:order_note_excel'))
        self.assertEqual(r.status_code, 200)
        self.assertIn('spreadsheetml', r['Content-Type'])

    def test_excel_export_empty_sheet_no_file(self):
        r = self.client.get(reverse('accounting:order_note_excel'))
        self.assertNotEqual(r.get('Content-Type', ''),
                            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
