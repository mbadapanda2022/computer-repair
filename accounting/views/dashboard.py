# accounting/views/dashboard.py

import csv
import logging
from decimal import Decimal
from datetime import date, timedelta, datetime

from django.shortcuts import render, redirect
from django.http import HttpResponse, JsonResponse
from django.db.models import Sum, Count, Q, F, Avg, Max, Min
from django.contrib import messages
from django.utils import timezone
from django.core.paginator import Paginator

from ..models import (
    Invoice, InvoiceItem, Purchase, Contact, Product,
    RepairJob, Payment, CompanyProfile, Notification
)
from .utils import is_htmx
from ..decorators import handle_errors

logger = logging.getLogger(__name__)

# Excel export
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:
    openpyxl = None


# ============================================================
# DASHBOARD (Staff) — COMPLETE
# ============================================================

@handle_errors(default_redirect='accounting:dashboard')
def dashboard(request):
    """Complete staff dashboard with stats, charts, and recent activity."""
    today = date.today()
    month_start = today.replace(day=1)
    week_start = today - timedelta(days=7)
    
    try:
        # ===== 1. STATS CARDS =====
        today_sales = Invoice.objects.filter(date=today).aggregate(
            total=Sum('grand_total')
        )['total'] or Decimal('0')
        
        month_sales = Invoice.objects.filter(date__gte=month_start).aggregate(
            total=Sum('grand_total')
        )['total'] or Decimal('0')
        
        pending_repairs = RepairJob.objects.filter(
            status__in=['pending', 'diagnosis', 'repairing', 'ready']
        ).count()
        
        low_stock = Product.objects.filter(
            is_service=False,
            current_stock__lte=F('low_stock_threshold')
        ).count()
        
        total_customers = Contact.objects.filter(
            contact_type__in=['customer', 'both']
        ).count()
        
        total_vendors = Contact.objects.filter(
            contact_type__in=['vendor', 'both']
        ).count()
        
        total_products = Product.objects.filter(is_active=True).count()
        
        pending_invoices = Invoice.objects.filter(
            payment_status__in=['unpaid', 'partial']
        ).count()
        
        total_invoices = Invoice.objects.count()
        
        # ===== 2. RECENT TRANSACTIONS =====
        recent_invoices = Invoice.objects.select_related('customer').order_by('-date')[:10]
        recent_repairs = RepairJob.objects.select_related('customer').order_by('-date_in')[:10]
        recent_payments = Payment.objects.select_related('contact').order_by('-date')[:10]
        
        # ===== 3. QUICK STATS =====
        total_payments_today = Payment.objects.filter(
            date=today,
            direction='received'
        ).aggregate(total=Sum('amount'))['total'] or Decimal('0')
        
        total_paid_amount = Payment.objects.filter(
            direction='received'
        ).aggregate(total=Sum('amount'))['total'] or Decimal('0')
        
        total_due_amount = Invoice.objects.aggregate(
            total=Sum('balance_due')
        )['total'] or Decimal('0')
        
        # ===== 4. UNREAD NOTIFICATIONS =====
        unread_notifications = Notification.objects.filter(
            recipient=request.user,
            is_read=False
        ).count()
        
        context = {
            # Stats
            'today_sales': today_sales,
            'month_sales': month_sales,
            'pending_repairs': pending_repairs,
            'low_stock': low_stock,
            'total_customers': total_customers,
            'total_vendors': total_vendors,
            'total_products': total_products,
            'pending_invoices': pending_invoices,
            'total_invoices': total_invoices,
            'total_payments_today': total_payments_today,
            'total_paid_amount': total_paid_amount,
            'total_due_amount': total_due_amount,
            'unread_notifications': unread_notifications,
            # Recent
            'recent_invoices': recent_invoices,
            'recent_repairs': recent_repairs,
            'recent_payments': recent_payments,
            # Date info
            'today': today,
            'month_start': month_start,
            'week_start': week_start,
        }
        
        # HTMX partial request (refresh stats)
        if is_htmx(request):
            return render(request, 'dashboard/partials/stats_cards.html', context)
        
        return render(request, 'dashboard/dashboard.html', context)
        
    except Exception as e:
        logger.error(f"Dashboard error: {e}")
        messages.error(request, "Unable to load dashboard.")
        return render(request, 'dashboard/dashboard.html', {
            'today_sales': Decimal('0'),
            'month_sales': Decimal('0'),
            'pending_repairs': 0,
            'low_stock': 0,
            'total_customers': 0,
            'total_vendors': 0,
            'total_products': 0,
            'pending_invoices': 0,
            'total_invoices': 0,
            'total_payments_today': Decimal('0'),
            'total_paid_amount': Decimal('0'),
            'total_due_amount': Decimal('0'),
            'unread_notifications': 0,
            'recent_invoices': [],
            'recent_repairs': [],
            'recent_payments': [],
        })


# ============================================================
# DASHBOARD STATS (HTMX / JSON) — CHARTS DATA
# ============================================================

@handle_errors(default_redirect='accounting:dashboard')
def dashboard_stats(request):
    """Return chart data as JSON for dashboard charts."""
    try:
        today = date.today()
        sales_data = []
        purchase_data = []
        repair_data = []
        labels = []

        # Last 7 days
        for i in range(6, -1, -1):
            d = today - timedelta(days=i)
            labels.append(d.strftime('%d %b'))
            
            sales_data.append(
                float(Invoice.objects.filter(date=d).aggregate(
                    total=Sum('grand_total')
                )['total'] or Decimal('0'))
            )
            
            purchase_data.append(
                float(Purchase.objects.filter(date=d).aggregate(
                    total=Sum('grand_total')
                )['total'] or Decimal('0'))
            )
            
            repair_data.append(
                float(RepairJob.objects.filter(date_in=d).count())
            )

        # Top products
        top_products = (
            InvoiceItem.objects
            .values('product__name')
            .annotate(qty=Sum('quantity'))
            .order_by('-qty')[:5]
        )
        
        top_products_labels = [p['product__name'] or 'Unknown' for p in top_products]
        top_products_data = [float(p['qty']) for p in top_products]

        # Payment methods distribution
        payment_methods = Payment.objects.values('method').annotate(
            total=Sum('amount')
        ).order_by('-total')
        
        payment_labels = [dict(Payment.METHOD_CHOICES).get(p['method'], p['method']) for p in payment_methods]
        payment_data = [float(p['total'] or 0) for p in payment_methods]

        # Repair status distribution
        repair_statuses = RepairJob.objects.values('status').annotate(
            count=Count('id')
        ).order_by('status')
        
        status_labels = [dict(RepairJob.STATUS_CHOICES).get(s['status'], s['status']) for s in repair_statuses]
        status_data = [s['count'] for s in repair_statuses]

        return JsonResponse({
            'labels': labels,
            'sales': sales_data,
            'purchases': purchase_data,
            'repairs': repair_data,
            'top_products_labels': top_products_labels,
            'top_products_data': top_products_data,
            'payment_labels': payment_labels,
            'payment_data': payment_data,
            'status_labels': status_labels,
            'status_data': status_data,
        })
        
    except Exception as e:
        logger.error(f"Dashboard stats error: {e}")
        return JsonResponse(
            {'error': 'Unable to load statistics. Please try again later.'},
            status=500
        )


# ============================================================
# DASHBOARD REFRESH STATS (HTMX)
# ============================================================

@handle_errors(default_redirect='accounting:dashboard')
def refresh_stats(request):
    """Refresh dashboard stats cards via HTMX."""
    today = date.today()
    month_start = today.replace(day=1)
    
    context = {
        'today_sales': Invoice.objects.filter(date=today).aggregate(
            total=Sum('grand_total')
        )['total'] or Decimal('0'),
        'month_sales': Invoice.objects.filter(date__gte=month_start).aggregate(
            total=Sum('grand_total')
        )['total'] or Decimal('0'),
        'pending_repairs': RepairJob.objects.filter(
            status__in=['pending', 'diagnosis', 'repairing', 'ready']
        ).count(),
        'low_stock': Product.objects.filter(
            is_service=False,
            current_stock__lte=F('low_stock_threshold')
        ).count(),
        'total_customers': Contact.objects.filter(
            contact_type__in=['customer', 'both']
        ).count(),
        'total_vendors': Contact.objects.filter(
            contact_type__in=['vendor', 'both']
        ).count(),
        'total_products': Product.objects.filter(is_active=True).count(),
        'pending_invoices': Invoice.objects.filter(
            payment_status__in=['unpaid', 'partial']
        ).count(),
        'total_payments_today': Payment.objects.filter(
            date=today,
            direction='received'
        ).aggregate(total=Sum('amount'))['total'] or Decimal('0'),
        'total_due_amount': Invoice.objects.aggregate(
            total=Sum('balance_due')
        )['total'] or Decimal('0'),
        'unread_notifications': Notification.objects.filter(
            recipient=request.user,
            is_read=False
        ).count(),
    }
    return render(request, 'dashboard/partials/stats_cards.html', context)


# ============================================================
# RECENT TRANSACTIONS (HTMX)
# ============================================================

@handle_errors(default_redirect='accounting:dashboard')
def recent_transactions(request):
    """Load recent transactions via HTMX."""
    try:
        # Get recent invoices
        recent_invoices = Invoice.objects.select_related('customer').order_by('-date')[:10]
        recent_repairs = RepairJob.objects.select_related('customer').order_by('-date_in')[:10]
        recent_payments = Payment.objects.select_related('contact').order_by('-date')[:10]
        
        context = {
            'recent_invoices': recent_invoices,
            'recent_repairs': recent_repairs,
            'recent_payments': recent_payments,
        }
        return render(request, 'dashboard/partials/recent_transactions.html', context)
        
    except Exception as e:
        logger.error(f"Recent transactions error: {e}")
        return render(request, 'dashboard/partials/recent_transactions.html', {
            'recent_invoices': [],
            'recent_repairs': [],
            'recent_payments': [],
        })


# ============================================================
# DASHBOARD EXPORT (Excel)
# ============================================================

@handle_errors(default_redirect='accounting:dashboard')
def dashboard_export(request):
    """Export dashboard summary to Excel."""
    if openpyxl is None:
        messages.error(request, "Openpyxl is not installed. Cannot export Excel.")
        return redirect('accounting:dashboard')
    
    try:
        today = date.today()
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Dashboard Summary"
        
        # Styles
        header_font = Font(bold=True, color="FFFFFF", size=12)
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        money_format = '#,##0.00'
        
        company = CompanyProfile.get_instance()
        
        # Title
        ws.merge_cells('A1:B1')
        ws['A1'] = company.name
        ws['A1'].font = Font(bold=True, size=14)
        ws.merge_cells('A2:B2')
        ws['A2'] = f"Dashboard Summary - {today.strftime('%d-%m-%Y')}"
        ws['A2'].font = Font(bold=True, size=12)
        
        # Stats
        row = 4
        stats = [
            ('Today\'s Sales', Invoice.objects.filter(date=today).aggregate(total=Sum('grand_total'))['total'] or Decimal('0')),
            ('Month Sales', Invoice.objects.filter(date__gte=today.replace(day=1)).aggregate(total=Sum('grand_total'))['total'] or Decimal('0')),
            ('Pending Repairs', RepairJob.objects.filter(status__in=['pending', 'diagnosis', 'repairing', 'ready']).count()),
            ('Low Stock Items', Product.objects.filter(is_service=False, current_stock__lte=F('low_stock_threshold')).count()),
            ('Total Customers', Contact.objects.filter(contact_type__in=['customer', 'both']).count()),
            ('Total Vendors', Contact.objects.filter(contact_type__in=['vendor', 'both']).count()),
            ('Total Products', Product.objects.filter(is_active=True).count()),
            ('Pending Invoices', Invoice.objects.filter(payment_status__in=['unpaid', 'partial']).count()),
            ('Total Paid Amount', Payment.objects.filter(direction='received').aggregate(total=Sum('amount'))['total'] or Decimal('0')),
            ('Total Due Amount', Invoice.objects.aggregate(total=Sum('balance_due'))['total'] or Decimal('0')),
        ]
        
        for label, value in stats:
            ws.cell(row=row, column=1, value=label)
            if isinstance(value, Decimal):
                ws.cell(row=row, column=2, value=float(value))
                ws.cell(row=row, column=2).number_format = money_format
            else:
                ws.cell(row=row, column=2, value=value)
            for col in range(1, 3):
                ws.cell(row=row, column=col).border = border
            row += 1
        
        # Column widths
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 20
        
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="dashboard_summary_{today.strftime("%Y%m%d")}.xlsx"'
        wb.save(response)
        return response
        
    except Exception as e:
        logger.error(f"Dashboard export error: {e}")
        messages.error(request, "Failed to export dashboard.")
        return redirect('accounting:dashboard')