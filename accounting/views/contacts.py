# accounting/views/contacts.py
import json
import logging
from decimal import Decimal

from django.contrib import messages
from django.core.paginator import EmptyPage, PageNotAnInteger, Paginator
from django.db.models import Q, Sum
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.csrf import csrf_protect
from django.views.decorators.http import require_http_methods

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side, numbers
from openpyxl.utils import get_column_letter

from ..decorators import handle_errors
from ..forms import ContactForm
from ..models import Contact, Invoice, LedgerLine, Payment, RepairJob
from .utils import htmx_field_error_response, is_htmx, redirect_to_staff

logger = logging.getLogger(__name__)


# ============================================================
# HELPERS
# ============================================================

def _has_real_transactions(contact):
    """
    Return True if the contact has any real transaction.

    Opening-balance ledger lines are excluded so that a contact whose
    only activity is an opening balance can still be deleted.
    """
    return bool(
        contact.sales_invoices.exists()
        or contact.purchases.exists()
        or contact.repair_jobs.exists()
        or contact.payments.exists()
        or contact.ledger_lines.exclude(
            ledger_entry__entry_type='opening'
        ).exists()
    )


def annotate_contact_balances(contacts):
    """
    Annotate each Contact with receivable, payable, and net position.

    Uses ACCOUNT CODES as the single source of truth so contact list
    numbers MATCH the statement view exactly:

      Receivable side (they owe us):
        1011 — Customer Receivable
        1012 — Advance from Customer (reduces receivable)

      Payable side (we owe them):
        2011 — Vendor Payable
        1014 — Advance to Vendor (reduces payable)

    Why account codes and not subledger_type?
      subledger_type is only set for 1011 and 2011 (not for advances).
      Using account codes ensures advances are correctly included,
      so contact list and statement never disagree.
    """
    if not contacts:
        return

    contact_ids = [c.id for c in contacts]

    RECV_CODES = ['1011', '1012']
    PAY_CODES = ['2011', '1014']

    # ── Receivable side ──
    recv_rows = (
        LedgerLine.objects
        .filter(
            contact__in=contact_ids,
            account__code__in=RECV_CODES,
            ledger_entry__is_deleted=False,   
        )
        .values('contact')
        .annotate(dr=Sum('debit'), cr=Sum('credit'))
    )
    recv_map = {
        r['contact']: (r['dr'] or Decimal('0')) - (r['cr'] or Decimal('0'))
        for r in recv_rows
    }

    # ── Payable side ──
    pay_rows = (
        LedgerLine.objects
        .filter(
            contact__in=contact_ids,
            account__code__in=PAY_CODES,
            ledger_entry__is_deleted=False,   
        )
        .values('contact')
        .annotate(dr=Sum('debit'), cr=Sum('credit'))
    )
    pay_map = {
        p['contact']: (p['cr'] or Decimal('0')) - (p['dr'] or Decimal('0'))
        for p in pay_rows
    }

    # ── Attach to each contact ──
    for contact in contacts:
        receivable = recv_map.get(contact.id, Decimal('0')).quantize(Decimal('0.01'))
        payable = pay_map.get(contact.id, Decimal('0')).quantize(Decimal('0.01'))

        contact.receivable_balance = receivable
        contact.payable_balance = payable

        # Net position
        #   customer  → receivable − payable
        #   vendor    → payable (only)
        #   both      → receivable − payable (net exposure)
        if contact.contact_type == 'vendor':
            contact.current_balance = payable
        else:
            contact.current_balance = (receivable - payable).quantize(Decimal('0.01'))


def get_paginated_contacts_context(request, queryset=None, filters=None):
    """
    Filter, paginate, and annotate contacts for list views.

    Default sort is newest-first so a freshly created contact always
    appears at the top of the list without any manual search.

    `filters` may be provided explicitly (e.g. from POST hidden inputs);
    otherwise query params are read from the request.
    """
    if filters is None:
        filters = {
            'search': request.GET.get('search', '').strip(),
            'contact_type': request.GET.get('contact_type', ''),
            'sort': request.GET.get('sort', 'recent'),
            'page': request.GET.get('page', 1),
        }

    search = (filters.get('search') or '').strip()
    contact_type = filters.get('contact_type') or ''
    sort = filters.get('sort') or 'recent'

    try:
        page_number = int(filters.get('page') or 1)
    except (ValueError, TypeError):
        page_number = 1

    if queryset is None:
        queryset = Contact.objects.all()

    sort_map = {
        'recent': ('-created_at', '-id'),
        'oldest': ('created_at', 'id'),
        'name': ('name',),
        'name_desc': ('-name',),
    }
    queryset = queryset.order_by(*sort_map.get(sort, ('-created_at', '-id')))

    if search:
        queryset = queryset.filter(
            Q(name__icontains=search) |
            Q(phone__icontains=search) |
            Q(email__icontains=search) |
            Q(company_name__icontains=search) |
            Q(gstin__icontains=search)
        )
    if contact_type:
        queryset = queryset.filter(contact_type=contact_type)

    paginator = Paginator(queryset, 25)
    try:
        page_obj = paginator.page(page_number)
    except PageNotAnInteger:
        page_obj = paginator.page(1)
    except EmptyPage:
        page_obj = paginator.page(paginator.num_pages)

    annotate_contact_balances(page_obj.object_list)

    return {
        'page_obj': page_obj,
        'contacts': page_obj.object_list,
        'search': search,
        'contact_type_filter': contact_type,
        'sort': sort,
    }


def _form_filter_context(request):
    """Build the filter context that the form template receives."""
    return {
        'filter_search': request.GET.get('search', ''),
        'filter_contact_type': request.GET.get('contact_type', ''),
        'filter_sort': request.GET.get('sort', 'recent'),
        'filter_page': request.GET.get('page', 1),
    }


# ============================================================
# FIELD VALIDATION (HTMX)
# ============================================================

@require_http_methods(["GET"])
def validate_contact_field(request):
    """Real-time validation for Contact fields."""
    field_name = request.GET.get('field')
    if not field_name:
        return HttpResponse("")

    value = request.GET.get(field_name, '').strip()
    contact_id = request.GET.get('contact_id')

    try:
        instance = Contact.objects.get(pk=contact_id) if contact_id else None
    except (Contact.DoesNotExist, ValueError):
        instance = None

    try:
        form = ContactForm(data={field_name: value}, instance=instance)
        form.full_clean()
        errors = form.errors.get(field_name, [])
    except Exception:
        # Never let a validator crash the page — return generic error
        errors = ['Unable to validate this field. Please check the value.']

    return htmx_field_error_response(field_name, errors)


# ============================================================
# CONTACT LIST
# ============================================================

@handle_errors(default_redirect='accounting:contact_list')
def contact_list(request):
    context = get_paginated_contacts_context(request)
    if is_htmx(request):
        return render(request, 'contacts/partials/contact_table.html', context)
    return render(request, 'contacts/contact_list.html', context)


# ============================================================
# EXCEL EXPORT
# ============================================================

@require_http_methods(["GET"])
def export_contacts_excel(request):
    """Export all contacts to a professional Excel file."""
    contacts = list(Contact.objects.all().order_by('-created_at', '-id'))
    annotate_contact_balances(contacts)

    wb = Workbook()
    ws = wb.active
    ws.title = "Contacts"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    title_font = Font(bold=True, size=14, color="1F4E78")
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'),
    )
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)
    right_align = Alignment(horizontal='right', vertical='center')
    money_format = numbers.FORMAT_NUMBER_COMMA_SEPARATED1

    TOTAL_COLS = 13
    last_col = get_column_letter(TOTAL_COLS)

    ws.merge_cells(f'A1:{last_col}1')
    title_cell = ws.cell(row=1, column=1, value="Contact List")
    title_cell.font = title_font
    title_cell.alignment = center_align
    ws.row_dimensions[1].height = 30

    headers = [
        'ID', 'Type', 'Name', 'Company', 'Phone', 'Email',
        'GSTIN', 'State', 'Opening',
        'Receivable', 'Payable', 'Net Position', 'Address',
    ]
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = center_align
    ws.row_dimensions[2].height = 25

    for idx, contact in enumerate(contacts, start=3):
        row_fill = PatternFill(
            start_color="F2F6FC" if idx % 2 == 0 else "FFFFFF",
            end_color="F2F6FC" if idx % 2 == 0 else "FFFFFF",
            fill_type="solid",
        )
        row_data = [
            contact.id,
            contact.get_contact_type_display(),
            contact.name,
            contact.company_name or '',
            contact.phone or '',
            contact.email or '',
            contact.gstin or '',
            contact.state or '',
            float(contact.opening_balance or 0),
            float(contact.receivable_balance or 0),
            float(contact.payable_balance or 0),
            float(contact.current_balance or 0),
            contact.address or '',
        ]

        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=idx, column=col_idx, value=value)
            cell.border = thin_border
            cell.fill = row_fill
            if col_idx in (1, 2, 5):
                cell.alignment = center_align
            elif col_idx in (9, 10, 11, 12):
                cell.alignment = right_align
                cell.number_format = money_format
            else:
                cell.alignment = left_align

    for col in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = max(min(max_length + 3, 45), 12)

    ws.freeze_panes = 'A3'

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="contacts.xlsx"'
    wb.save(response)
    return response


# ============================================================
# CONTACT CREATE
# ============================================================

@csrf_protect
@handle_errors(
    default_redirect='accounting:contact_list',
    htmx_template='contacts/contact_form.html',
)
def contact_create(request):
    if request.method == 'POST':
        form = ContactForm(request.POST)
        if form.is_valid():
            contact = form.save()
            logger.info(
                "Contact '%s' created by %s",
                contact.name, request.user.username,
            )

            if is_htmx(request):
                # After a successful create, reset filters so the new
                # contact is always visible at the top of a fresh list.
                filters = {
                    'search': '',
                    'contact_type': '',
                    'sort': 'recent',
                    'page': 1,
                }
                context = get_paginated_contacts_context(request, filters=filters)
                response = render(request, 'contacts/partials/contact_table.html', context)
                response['HX-Trigger'] = json.dumps({
                    'showToast': {
                        'level': 'success',
                        'message': f'Contact "{contact.name}" created.',
                    },
                    'closeModal': '',
                })
                return response

            messages.success(request, f'Contact "{contact.name}" created successfully.')
            return redirect_to_staff('contact_list')

        # Invalid form
        if is_htmx(request):
            response = render(
                request, 'contacts/contact_form.html',
                {'form': form, **_form_filter_context(request)},
            )
            response['HX-Retarget'] = '#mainModalContent'
            return response

        return render(
            request, 'contacts/contact_form.html',
            {'form': form, **_form_filter_context(request)},
        )

    # GET
    form = ContactForm()
    return render(request, 'contacts/contact_form.html', {
        'form': form,
        **_form_filter_context(request),
    })


# ============================================================
# CONTACT UPDATE
# ============================================================

@csrf_protect
@handle_errors(
    default_redirect='accounting:contact_list',
    htmx_template='contacts/contact_form.html',
)
def contact_update(request, pk):
    contact = get_object_or_404(Contact, pk=pk)

    if request.method == 'POST':
        form = ContactForm(request.POST, instance=contact)
        if form.is_valid():
            contact = form.save()
            logger.info(
                "Contact '%s' updated by %s",
                contact.name, request.user.username,
            )

            if is_htmx(request):
                # After update, preserve the user's current filter/view
                # so they stay in the same context they were browsing.
                filters = {
                    'search': request.POST.get('_filter_search', ''),
                    'contact_type': request.POST.get('_filter_contact_type', ''),
                    'sort': request.POST.get('_filter_sort', 'recent'),
                    'page': request.POST.get('_filter_page', 1),
                }
                context = get_paginated_contacts_context(request, filters=filters)
                response = render(request, 'contacts/partials/contact_table.html', context)
                response['HX-Trigger'] = json.dumps({
                    'showToast': {
                        'level': 'success',
                        'message': f'Contact "{contact.name}" updated.',
                    },
                    'closeModal': '',
                })
                return response

            messages.success(request, f'Contact "{contact.name}" updated successfully.')
            return redirect_to_staff('contact_list')

        if is_htmx(request):
            response = render(
                request, 'contacts/contact_form.html',
                {'form': form, 'contact': contact, **_form_filter_context(request)},
            )
            response['HX-Retarget'] = '#mainModalContent'
            return response

        return render(
            request, 'contacts/contact_form.html',
            {'form': form, 'contact': contact, **_form_filter_context(request)},
        )

    form = ContactForm(instance=contact)
    return render(request, 'contacts/contact_form.html', {
        'form': form,
        'contact': contact,
        **_form_filter_context(request),
    })


# ============================================================
# CONTACT DELETE
# ============================================================

@csrf_protect
@require_http_methods(["DELETE"])
@handle_errors(default_redirect='accounting:contact_list')
def contact_delete(request, pk):
    contact = get_object_or_404(Contact, pk=pk)

    if _has_real_transactions(contact):
        response = HttpResponse(
            "Cannot delete contact with existing transactions.", status=400,
        )
        response['HX-Trigger'] = json.dumps({
            'showToast': {
                'level': 'danger',
                'message': f'"{contact.name}" has transactions and cannot be deleted.',
            }
        })
        return response

    name = contact.name
    contact.delete()  # soft delete

    context = get_paginated_contacts_context(request)
    response = render(request, 'contacts/partials/contact_table.html', context)
    response['HX-Trigger'] = json.dumps({
        'showToast': {'level': 'success', 'message': f'Contact "{name}" deleted.'}
    })
    return response


# ============================================================
# CONTACT DETAIL MODAL
# ============================================================

def contact_detail_modal(request, pk):
    contact = get_object_or_404(Contact, pk=pk)

    invoices = Invoice.objects.filter(customer=contact)
    total_invoiced = invoices.aggregate(t=Sum('grand_total'))['t'] or Decimal('0')
    total_due = invoices.aggregate(t=Sum('balance_due'))['t'] or Decimal('0')

    purchases = contact.purchases.all()
    total_purchased = purchases.aggregate(t=Sum('grand_total'))['t'] or Decimal('0')

    repairs = RepairJob.objects.filter(customer=contact)

    payments_received = Payment.objects.filter(contact=contact, direction='received')
    payments_paid = Payment.objects.filter(contact=contact, direction='paid')
    total_received = payments_received.aggregate(t=Sum('amount'))['t'] or Decimal('0')
    total_paid = payments_paid.aggregate(t=Sum('amount'))['t'] or Decimal('0')

    context = {
        'contact': contact,
        'invoice_count': invoices.count(),
        'total_invoiced': total_invoiced,
        'total_due': total_due,
        'purchase_count': purchases.count(),
        'total_purchased': total_purchased,
        'repair_count': repairs.count(),
        'total_received': total_received,
        'total_paid': total_paid,
    }
    return render(request, 'contacts/partials/contact_detail_modal.html', context)


# ============================================================
# QUICK REDIRECTS
# ============================================================

def quick_journal(request, pk):
    contact = get_object_or_404(Contact, pk=pk)
    return redirect_to_staff('journal_create_for_contact', contact_id=contact.pk)


def quick_statement(request, pk):
    contact = get_object_or_404(Contact, pk=pk)
    if contact.contact_type in ('customer', 'both'):
        return redirect_to_staff('customer_statement', contact_id=contact.pk)
    return redirect_to_staff('vendor_statement', contact_id=contact.pk)


# ============================================================
# BULK DELETE
# ============================================================

@csrf_protect
@require_http_methods(["POST"])
@handle_errors(default_redirect='accounting:contact_list')
def bulk_delete_contacts(request):
    ids = request.POST.getlist('ids')
    if not ids:
        return JsonResponse(
            {'status': 'error', 'message': 'No contacts selected.'}, status=400,
        )

    contacts = Contact.objects.filter(pk__in=ids)

    blocking = []
    for contact in contacts:
        if _has_real_transactions(contact):
            blocking.append(contact.name)

    if blocking:
        return JsonResponse({
            'status': 'error',
            'message': f'Cannot delete: {", ".join(blocking)} (has transactions).',
        }, status=400)

    count = contacts.count()
    contacts.delete()  # soft delete all
    return JsonResponse({'status': 'ok', 'deleted': count})


# ============================================================
# CONTACT SEARCH (JSON autocomplete)
# ============================================================

@require_http_methods(["GET"])
def contact_search(request):
    q = request.GET.get('q', '').strip()
    results = []

    if len(q) >= 1:
        contacts = Contact.objects.filter(
            Q(name__icontains=q) |
            Q(phone__icontains=q) |
            Q(email__icontains=q) |
            Q(company_name__icontains=q)
        )[:20]

        results = [{
            'id': c.id,
            'name': c.name,
            'phone': c.phone or '',
            'email': c.email or '',
            'type': c.contact_type,
            'type_display': c.get_contact_type_display(),
            'label': f"{c.name} ({c.phone or 'No phone'})",
        } for c in contacts]

    return JsonResponse({'results': results})