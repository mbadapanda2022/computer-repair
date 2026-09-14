# accounting/views/bank.py

import json
import logging
from decimal import Decimal

from django.shortcuts import render, redirect, get_object_or_404
from django.http import HttpResponse, JsonResponse
from django.contrib import messages
from django.db import transaction
from django.db.models import Q, Sum
from django.core.paginator import Paginator, EmptyPage, PageNotAnInteger
from django.template.loader import render_to_string
from django.urls import reverse
from django.views.decorators.csrf import csrf_protect
from django.views.decorators.http import require_http_methods
from django.utils import timezone

from ..models import BankAccount, BankTransaction, CompanyProfile
from ..forms import BankAccountForm, BankTransactionForm
from .utils import is_htmx, htmx_response, redirect_to_staff, toast_only_response
from ..decorators import handle_errors

logger = logging.getLogger(__name__)

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    openpyxl = None


# =============================================================
# HELPER: RECALCULATE ACCOUNT BALANCE
# =============================================================
def recalculate_account_balance(account):
    """Recalculate and update current_balance for a bank account."""
    total_deposits = account.transactions.filter(
        transaction_type='deposit'
    ).aggregate(total=Sum('amount'))['total'] or Decimal('0')

    total_withdrawals = account.transactions.filter(
        transaction_type='withdrawal'
    ).aggregate(total=Sum('amount'))['total'] or Decimal('0')

    account.current_balance = (
        account.opening_balance + total_deposits - total_withdrawals
    ).quantize(Decimal('0.01'))
    account.save(update_fields=['current_balance'])
    return account.current_balance


# =============================================================
# BANK ACCOUNT LIST
# =============================================================
def bank_account_list(request):
    """Bank account list with summary cards."""
    search = request.GET.get('search', '').strip()

    accounts = BankAccount.objects.all().order_by('name')
    if search:
        accounts = accounts.filter(
            Q(name__icontains=search) |
            Q(bank_name__icontains=search) |
            Q(account_number__icontains=search)
        )

    # Summary totals
    total_opening = accounts.aggregate(t=Sum('opening_balance'))['t'] or Decimal('0')
    total_current = accounts.aggregate(t=Sum('current_balance'))['t'] or Decimal('0')

    context = {
        'accounts': accounts,
        'search': search,
        'total_opening': total_opening,
        'total_current': total_current,
    }

    if is_htmx(request):
        return render(request, 'bank/partials/bank_account_table.html', context)
    return render(request, 'bank/bank_account_list.html', context)


# =============================================================
# BANK ACCOUNT CREATE
# =============================================================
@csrf_protect
@handle_errors(default_redirect='accounting:bank_account_list',
               htmx_template='bank/partials/bank_account_form.html')
def bank_account_add(request):
    if request.method == 'POST':
        form = BankAccountForm(request.POST)
        if form.is_valid():
            with transaction.atomic():
                account = form.save()
                account.current_balance = account.opening_balance
                account.save(update_fields=['current_balance'])

            logger.info(f"Bank account created: {account.name} by {request.user.username}")

            if is_htmx(request):
                # Close modal + reload the account list section
                response = HttpResponse()
                response['HX-Trigger'] = json.dumps({
                    'closeModal': '',
                    'showToast': {
                        'level': 'success',
                        'message': f'Bank account "{account.name}" created.'
                    },
                    'reloadBankAccounts': ''
                })
                return response

            messages.success(request, f'Bank account "{account.name}" created successfully.')
            return redirect_to_staff('bank_account_list')
        else:
            if is_htmx(request):
                return render(request, 'bank/partials/bank_account_form.html', {'form': form})
    else:
        form = BankAccountForm()

    context = {'form': form}
    if is_htmx(request):
        return render(request, 'bank/partials/bank_account_form.html', context)
    return render(request, 'bank/bank_account_form.html', context)


# =============================================================
# BANK ACCOUNT EDIT
# =============================================================
@csrf_protect
@handle_errors(default_redirect='accounting:bank_account_list',
               htmx_template='bank/partials/bank_account_form.html')
def bank_account_edit(request, pk):
    account = get_object_or_404(BankAccount, pk=pk)

    if request.method == 'POST':
        form = BankAccountForm(request.POST, instance=account)
        if form.is_valid():
            with transaction.atomic():
                form.save()
                recalculate_account_balance(account)

            logger.info(f"Bank account updated: {account.name} by {request.user.username}")

            if is_htmx(request):
                response = HttpResponse()
                response['HX-Trigger'] = json.dumps({
                    'closeModal': '',
                    'showToast': {
                        'level': 'success',
                        'message': f'Bank account "{account.name}" updated.'
                    },
                    'reloadBankAccounts': ''
                })
                return response

            messages.success(request, f'Bank account "{account.name}" updated successfully.')
            return redirect_to_staff('bank_account_list')
        else:
            if is_htmx(request):
                return render(request, 'bank/partials/bank_account_form.html', {
                    'form': form, 'account': account
                })
    else:
        form = BankAccountForm(instance=account)

    context = {'form': form, 'account': account}
    if is_htmx(request):
        return render(request, 'bank/partials/bank_account_form.html', context)
    return render(request, 'bank/bank_account_form.html', context)


# =============================================================
# BANK ACCOUNT DELETE
# =============================================================
@csrf_protect
@require_http_methods(["DELETE"])
@handle_errors(default_redirect='accounting:bank_account_list')
def bank_account_delete(request, pk):
    account = get_object_or_404(BankAccount, pk=pk)

    if account.transactions.exists():
        response = HttpResponse(
            "Cannot delete account with existing transactions.",
            status=400
        )
        response['HX-Trigger'] = json.dumps({
            'showToast': {
                'level': 'danger',
                'message': f'Cannot delete "{account.name}" — it has transactions.'
            }
        })
        return response

    try:
        name = account.name
        account.delete()
        logger.info(f"Bank account deleted: {name} by {request.user.username}")

        # Return the whole table partial (with data re-queried)
        accounts = BankAccount.objects.all().order_by('name')
        total_opening = accounts.aggregate(t=Sum('opening_balance'))['t'] or Decimal('0')
        total_current = accounts.aggregate(t=Sum('current_balance'))['t'] or Decimal('0')

        response = render(request, 'bank/partials/bank_account_table.html', {
            'accounts': accounts,
            'total_opening': total_opening,
            'total_current': total_current,
        })
        response['HX-Trigger'] = json.dumps({
            'showToast': {'level': 'success', 'message': f'Bank account "{name}" deleted.'}
        })
        return response
    except Exception as e:
        logger.error(f"Error deleting bank account {pk}: {e}")
        return toast_only_response(
            {'level': 'danger', 'message': 'Failed to delete bank account.'},
            status=500
        )


# =============================================================
# BANK STATEMENT (with pagination, filters, HTMX)
# =============================================================
def bank_statement(request, pk):
    """Bank statement view with pagination, filters, print support."""
    account = get_object_or_404(BankAccount, pk=pk)

    date_from = request.GET.get('date_from', '').strip()
    date_to = request.GET.get('date_to', '').strip()
    txn_type = request.GET.get('txn_type', '')  # deposit | withdrawal
    search = request.GET.get('search', '').strip()
    page_number = request.GET.get('page', 1)
    is_print = request.GET.get('print') == '1'

    transactions_qs = account.transactions.all().order_by('date', 'id')

    if date_from:
        transactions_qs = transactions_qs.filter(date__gte=date_from)
    if date_to:
        transactions_qs = transactions_qs.filter(date__lte=date_to)
    if txn_type in ('deposit', 'withdrawal'):
        transactions_qs = transactions_qs.filter(transaction_type=txn_type)
    if search:
        transactions_qs = transactions_qs.filter(
            Q(description__icontains=search) |
            Q(reference__icontains=search)
        )

    # Opening balance with carry-forward
    opening_balance = account.opening_balance
    if date_from:
        prior = account.transactions.filter(date__lt=date_from)
        dep = prior.filter(transaction_type='deposit').aggregate(t=Sum('amount'))['t'] or Decimal('0')
        wd = prior.filter(transaction_type='withdrawal').aggregate(t=Sum('amount'))['t'] or Decimal('0')
        opening_balance = account.opening_balance + dep - wd

    # Build statement rows
    running_balance = opening_balance
    all_rows = []
    total_deposit = Decimal('0')
    total_withdrawal = Decimal('0')

    for txn in transactions_qs:
        if txn.transaction_type == 'deposit':
            running_balance += txn.amount
            deposit = txn.amount
            withdrawal = Decimal('0')
            total_deposit += deposit
        else:
            running_balance -= txn.amount
            deposit = Decimal('0')
            withdrawal = txn.amount
            total_withdrawal += withdrawal

        all_rows.append({
            'id': txn.id,
            'date': txn.date,
            'description': txn.description or txn.get_source_type_display(),
            'source_type': txn.source_type,
            'source_type_display': txn.get_source_type_display(),
            'deposit': deposit,
            'withdrawal': withdrawal,
            'balance': running_balance,
            'reference': txn.reference or '-',
            'reconciled': txn.reconciled,
            'payment_id': txn.payment_id,
        })

    closing_balance = running_balance

    company = CompanyProfile.get_instance()

    context = {
        'account': account,
        'company': company,
        'logo_exists': bool(company.logo and company.logo.name),
        'opening_balance': opening_balance,
        'closing_balance': closing_balance,
        'total_deposit': total_deposit,
        'total_withdrawal': total_withdrawal,
        'date_from': date_from,
        'date_to': date_to,
        'txn_type': txn_type,
        'search': search,
    }

    # Print mode: pass all rows (no pagination)
    if is_print:
        context['statement_lines'] = all_rows
        return render(request, 'bank/bank_statement_print.html', context)

    # Pagination (25 per page)
    paginator = Paginator(all_rows, 25)
    try:
        page_obj = paginator.page(page_number)
    except PageNotAnInteger:
        page_obj = paginator.page(1)
    except EmptyPage:
        page_obj = paginator.page(paginator.num_pages)

    context['statement_lines'] = page_obj.object_list
    context['page_obj'] = page_obj

    if is_htmx(request):
        return render(request, 'bank/partials/bank_statement_table.html', context)
    return render(request, 'bank/bank_statement.html', context)


# =============================================================
# BANK STATEMENT - Excel Export
# =============================================================
def bank_statement_excel(request, pk):
    """Export bank statement to Excel."""
    if openpyxl is None:
        return toast_only_response(
            {'level': 'danger', 'message': 'Openpyxl library is not installed.'},
            status=400
        )

    account = get_object_or_404(BankAccount, pk=pk)
    date_from = request.GET.get('date_from', '').strip()
    date_to = request.GET.get('date_to', '').strip()
    txn_type = request.GET.get('txn_type', '')
    search = request.GET.get('search', '').strip()

    transactions_qs = account.transactions.all().order_by('date', 'id')
    if date_from:
        transactions_qs = transactions_qs.filter(date__gte=date_from)
    if date_to:
        transactions_qs = transactions_qs.filter(date__lte=date_to)
    if txn_type in ('deposit', 'withdrawal'):
        transactions_qs = transactions_qs.filter(transaction_type=txn_type)
    if search:
        transactions_qs = transactions_qs.filter(
            Q(description__icontains=search) | Q(reference__icontains=search)
        )

    opening_balance = account.opening_balance
    if date_from:
        prior = account.transactions.filter(date__lt=date_from)
        dep = prior.filter(transaction_type='deposit').aggregate(t=Sum('amount'))['t'] or Decimal('0')
        wd = prior.filter(transaction_type='withdrawal').aggregate(t=Sum('amount'))['t'] or Decimal('0')
        opening_balance = account.opening_balance + dep - wd

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bank Statement"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    thin = Side(style='thin', color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left = Alignment(horizontal='left', vertical='center')
    right = Alignment(horizontal='right', vertical='center')
    money_fmt = '#,##0.00'

    ws.merge_cells('A1:F1')
    ws['A1'] = f"{account.name} — Bank Statement"
    ws['A1'].font = Font(bold=True, size=14, color="1F4E78")
    ws['A1'].alignment = center

    ws.merge_cells('A2:F2')
    ws['A2'] = f"Period: {date_from or 'Beginning'} to {date_to or 'Today'}"
    ws['A2'].alignment = center

    headers = ['Date', 'Description', 'Reference', 'Deposit (₹)', 'Withdrawal (₹)', 'Balance (₹)']
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=4, column=col, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center
        c.border = border

    row = 5
    # Opening balance row
    ws.cell(row=row, column=2, value='Opening Balance')
    ws.cell(row=row, column=6, value=float(opening_balance))
    ws.cell(row=row, column=6).number_format = money_fmt
    for col in range(1, 7):
        ws.cell(row=row, column=col).border = border
    row += 1

    running_balance = opening_balance
    total_dep = Decimal('0')
    total_wd = Decimal('0')

    for txn in transactions_qs:
        if txn.transaction_type == 'deposit':
            running_balance += txn.amount
            dep = txn.amount
            wd = Decimal('0')
            total_dep += dep
        else:
            running_balance -= txn.amount
            dep = Decimal('0')
            wd = txn.amount
            total_wd += wd

        ws.cell(row=row, column=1, value=txn.date.strftime('%d-%m-%Y'))
        ws.cell(row=row, column=2, value=txn.description or txn.get_source_type_display())
        ws.cell(row=row, column=3, value=txn.reference or '')
        ws.cell(row=row, column=4, value=float(dep) if dep else '').number_format = money_fmt
        ws.cell(row=row, column=5, value=float(wd) if wd else '').number_format = money_fmt
        ws.cell(row=row, column=6, value=float(running_balance)).number_format = money_fmt
        for col in range(1, 7):
            ws.cell(row=row, column=col).border = border
        row += 1

    # Total row
    ws.cell(row=row, column=2, value='Totals').font = Font(bold=True)
    ws.cell(row=row, column=4, value=float(total_dep)).number_format = money_fmt
    ws.cell(row=row, column=4).font = Font(bold=True)
    ws.cell(row=row, column=5, value=float(total_wd)).number_format = money_fmt
    ws.cell(row=row, column=5).font = Font(bold=True)
    ws.cell(row=row, column=6, value=float(running_balance)).number_format = money_fmt
    ws.cell(row=row, column=6).font = Font(bold=True)
    for col in range(1, 7):
        ws.cell(row=row, column=col).border = border

    widths = [14, 45, 20, 16, 16, 16]
    for idx, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    ws.freeze_panes = 'A5'

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    safe = account.name.replace(' ', '_').replace('/', '_')
    response['Content-Disposition'] = f'attachment; filename="statement_{safe}_{timezone.now().strftime("%Y%m%d")}.xlsx"'
    wb.save(response)
    return response


# =============================================================
# ALL TRANSACTIONS (cross-account list)
# =============================================================
def bank_transaction_list(request):
    """All transactions across all bank accounts."""
    search = request.GET.get('search', '').strip()
    account_id = request.GET.get('account', '')
    txn_type = request.GET.get('txn_type', '')
    date_from = request.GET.get('date_from', '').strip()
    date_to = request.GET.get('date_to', '').strip()
    page_number = request.GET.get('page', 1)

    qs = BankTransaction.objects.select_related(
        'bank_account', 'payment', 'invoice'
    ).order_by('-date', '-id')

    if search:
        qs = qs.filter(
            Q(description__icontains=search) |
            Q(reference__icontains=search) |
            Q(bank_account__name__icontains=search)
        )
    if account_id:
        qs = qs.filter(bank_account_id=account_id)
    if txn_type in ('deposit', 'withdrawal'):
        qs = qs.filter(transaction_type=txn_type)
    if date_from:
        qs = qs.filter(date__gte=date_from)
    if date_to:
        qs = qs.filter(date__lte=date_to)

    paginator = Paginator(qs, 25)
    try:
        page_obj = paginator.page(page_number)
    except PageNotAnInteger:
        page_obj = paginator.page(1)
    except EmptyPage:
        page_obj = paginator.page(paginator.num_pages)

    accounts = BankAccount.objects.filter(is_active=True).order_by('name')

    context = {
        'page_obj': page_obj,
        'transactions': page_obj.object_list,
        'accounts': accounts,
        'search': search,
        'account_id': account_id,
        'txn_type': txn_type,
        'date_from': date_from,
        'date_to': date_to,
    }

    if is_htmx(request):
        return render(request, 'bank/partials/bank_transaction_table.html', context)
    return render(request, 'bank/bank_transaction_list.html', context)


# =============================================================
# ADD BANK TRANSACTION
# =============================================================
@csrf_protect
@handle_errors(default_redirect='accounting:bank_account_list',
               htmx_template='bank/partials/bank_transaction_form.html')
def bank_transaction_add(request, account_pk):
    account = get_object_or_404(BankAccount, pk=account_pk)

    if request.method == 'POST':
        form = BankTransactionForm(request.POST)
        if form.is_valid():
            with transaction.atomic():
                txn = form.save(commit=False)
                txn.bank_account = account

                # Insufficient balance check
                if txn.transaction_type == 'withdrawal' and txn.amount > account.current_balance:
                    form.add_error(
                        'amount',
                        f'Insufficient balance. Available: ₹{account.current_balance}'
                    )
                    if is_htmx(request):
                        return render(request, 'bank/partials/bank_transaction_form.html', {
                            'form': form, 'account': account
                        })
                    messages.error(request, f'Insufficient balance. Available: ₹{account.current_balance}')
                    return redirect_to_staff('bank_statement', pk=account.pk)

                txn.save()
                logger.info(
                    f"Bank transaction added: {txn.get_transaction_type_display()} "
                    f"₹{txn.amount} for {account.name} by {request.user.username}"
                )
                recalculate_account_balance(account)

                if is_htmx(request):
                    # Redirect back to statement page
                    response = HttpResponse()
                    response['HX-Redirect'] = reverse('accounting:bank_statement', args=[account.pk])
                    response['HX-Trigger'] = json.dumps({
                        'closeModal': '',
                        'showToast': {
                            'level': 'success',
                            'message': 'Transaction recorded successfully.'
                        }
                    })
                    return response
                messages.success(request, 'Transaction recorded successfully.')
                return redirect_to_staff('bank_statement', pk=account.pk)
        else:
            if is_htmx(request):
                return render(request, 'bank/partials/bank_transaction_form.html', {
                    'form': form, 'account': account
                })
    else:
        form = BankTransactionForm(initial={
            'bank_account': account,
            'date': timezone.now().date(),
        })

    context = {'form': form, 'account': account}
    if is_htmx(request):
        return render(request, 'bank/partials/bank_transaction_form.html', context)
    return render(request, 'bank/bank_transaction_form.html', context)


# =============================================================
# EDIT BANK TRANSACTION
# =============================================================
@csrf_protect
@handle_errors(default_redirect='accounting:bank_account_list',
               htmx_template='bank/partials/bank_transaction_form.html')
def bank_transaction_edit(request, pk):
    txn = get_object_or_404(BankTransaction, pk=pk)
    account = txn.bank_account

    if request.method == 'POST':
        form = BankTransactionForm(request.POST, instance=txn)
        if form.is_valid():
            with transaction.atomic():
                updated = form.save()
                recalculate_account_balance(account)

            logger.info(f"Bank transaction #{pk} updated by {request.user.username}")

            if is_htmx(request):
                response = HttpResponse()
                response['HX-Redirect'] = reverse('accounting:bank_statement', args=[account.pk])
                response['HX-Trigger'] = json.dumps({
                    'closeModal': '',
                    'showToast': {'level': 'success', 'message': 'Transaction updated.'}
                })
                return response
            messages.success(request, 'Transaction updated.')
            return redirect_to_staff('bank_statement', pk=account.pk)
        else:
            if is_htmx(request):
                return render(request, 'bank/partials/bank_transaction_form.html', {
                    'form': form, 'account': account, 'txn': txn
                })
    else:
        form = BankTransactionForm(instance=txn)

    context = {'form': form, 'account': account, 'txn': txn}
    if is_htmx(request):
        return render(request, 'bank/partials/bank_transaction_form.html', context)
    return render(request, 'bank/bank_transaction_form.html', context)


# =============================================================
# DELETE BANK TRANSACTION
# =============================================================
@csrf_protect
@require_http_methods(["DELETE"])
@handle_errors(default_redirect='accounting:bank_account_list')
def bank_transaction_delete(request, pk):
    txn = get_object_or_404(BankTransaction, pk=pk)
    account = txn.bank_account

    try:
        txn.delete()
        recalculate_account_balance(account)
        logger.info(f"Bank transaction #{pk} deleted by {request.user.username}")

        response = HttpResponse()
        response['HX-Redirect'] = reverse('accounting:bank_statement', args=[account.pk])
        response['HX-Trigger'] = json.dumps({
            'showToast': {'level': 'success', 'message': 'Transaction deleted.'}
        })
        return response
    except Exception as e:
        logger.error(f"Error deleting bank transaction #{pk}: {e}")
        return toast_only_response(
            {'level': 'danger', 'message': 'Failed to delete transaction.'},
            status=500
        )