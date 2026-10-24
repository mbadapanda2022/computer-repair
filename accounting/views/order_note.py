# accounting/views/order_note.py
"""
Temporary supplier Order Note — an Excel-like scratch sheet kept in the
user's session (never in the DB). Export to .xlsx and send to the supplier
on WhatsApp so they can fill rates / availability themselves.
"""
import re
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.utils import timezone
from django.views.decorators.http import require_POST

from ..models import CompanyProfile
from .utils import toast_only_response

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    Workbook = None

SESSION_KEY = 'order_note'
TWO_PLACES = Decimal('0.01')


def _dec(value):
    """Loose numeric parse — this sheet has no validation, bad input = blank."""
    if value is None:
        return None
    s = str(value).replace(',', '').strip()
    if not s:
        return None
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def _money(value):
    if value is None:
        return None
    return value.quantize(TWO_PLACES, rounding=ROUND_HALF_UP)


def _compute_row(raw):
    """Normalise one sheet row and fill whichever price side is missing.

    price_src tells us which column the user typed last:
      'incl' -> derive excl from incl,  'excl'/default -> derive incl from excl.
    """
    qty = _dec(raw.get('qty')) or Decimal('0')
    tax = _dec(raw.get('tax'))
    if tax is None:
        tax = Decimal('18')
    excl = _dec(raw.get('price_excl'))
    incl = _dec(raw.get('price_incl'))
    factor = (Decimal('1') + tax / Decimal('100'))

    src = raw.get('price_src')
    if src == 'incl' and incl is not None:
        excl = _money(incl / factor)
    elif excl is not None:
        incl = _money(excl * factor)
    elif incl is not None:
        excl = _money(incl / factor)

    amount = _money(qty * incl) if incl is not None and qty else None
    line_excl = _money(qty * excl) if excl is not None and qty else None
    return {
        'date': (raw.get('date') or '').strip(),
        'product': (raw.get('product') or '').strip(),
        'qty': qty,
        'tax': tax,
        'price_excl': excl,
        'price_incl': incl,
        'price_src': 'incl' if src == 'incl' else 'excl',
        'amount': amount,
        'line_excl': line_excl,
    }


def _rows_from_post(request):
    def col(name):
        return request.POST.getlist(f'{name}[]')

    cols = {name: col(name) for name in
            ('date', 'product', 'qty', 'tax', 'price_excl', 'price_incl', 'price_src')}
    n = max((len(v) for v in cols.values()), default=0)
    rows = []
    for i in range(n):
        raw = {k: (v[i] if i < len(v) else '') for k, v in cols.items()}
        # date/tax come prefilled; a zero/junk qty isn't content either
        has_content = bool(raw['product'].strip()) or any(
            _dec(raw[k]) for k in ('qty', 'price_excl', 'price_incl'))
        if not has_content:
            continue
        rows.append(_compute_row(raw))
    return rows


def _totals(rows):
    subtotal_excl = sum((r['line_excl'] for r in rows if r['line_excl'] is not None), Decimal('0'))
    grand = sum((r['amount'] for r in rows if r['amount'] is not None), Decimal('0'))
    return {
        'qty': sum((r['qty'] for r in rows if r['qty'] is not None), Decimal('0')),
        'subtotal_excl': _money(subtotal_excl),
        'tax_total': _money(grand - subtotal_excl),
        'grand': _money(grand),
    }


def _get_sheet(request):
    return request.session.get(SESSION_KEY) or {'supplier': '', 'note': '', 'rows': []}


def _store_sheet(request):
    rows = _rows_from_post(request)
    request.session[SESSION_KEY] = {
        'supplier': request.POST.get('supplier', '').strip(),
        'note': request.POST.get('note', '').strip(),
        'rows': [{k: _cell(v) for k, v in r.items()} for r in rows],
    }
    return rows


def _cell(value):
    if value is None or value == Decimal('0'):
        return ''
    return str(value)


# ============================================================
# SHEET PAGE (create / edit / delete rows, all saved at once)
# ============================================================
def order_note(request):
    if request.method == 'POST':
        _store_sheet(request)
        return redirect('accounting:order_note')

    sheet = _get_sheet(request)
    rows = [_compute_row(r) for r in sheet['rows']]
    context = {
        'supplier': sheet.get('supplier', ''),
        'note': sheet.get('note', ''),
        'rows': rows,
        'totals': _totals(rows),
        'today': timezone.localdate().isoformat(),
        'company': CompanyProfile.get_instance(),
    }
    return render(request, 'order_note/sheet.html', context)


# ============================================================
# CLEAR (hard delete — nothing persists anyway)
# ============================================================
@require_POST
def order_note_clear(request):
    request.session.pop(SESSION_KEY, None)
    return redirect('accounting:order_note')


# ============================================================
# EXCEL EXPORT (POST saves the live grid first, then downloads)
# ============================================================
def order_note_excel(request):
    if request.method == 'POST':
        _store_sheet(request)
    sheet = _get_sheet(request)
    rows = [_compute_row(r) for r in sheet['rows']]
    if not rows:
        return toast_only_response(
            {'level': 'warning', 'message': 'Order note is empty — add some items first.'})
    if Workbook is None:  # pragma: no cover
        return toast_only_response(
            {'level': 'danger', 'message': 'Openpyxl library is not installed.'}, status=400)

    company = CompanyProfile.get_instance()
    supplier = sheet.get('supplier', '')

    wb = Workbook()
    ws = wb.active
    ws.title = 'Order Note'

    header_font = Font(bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
    thin = Side(style='medium', color='808080')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal='center', vertical='center')
    right_a = Alignment(horizontal='right', vertical='center')
    money_fmt = '#,##0.00'

    ws.merge_cells('A1:J1')
    ws['A1'] = f"{company.name or 'My Company'} — Order Note"
    ws['A1'].font = Font(bold=True, size=14, color='1F4E78')
    ws['A1'].alignment = center

    ws.merge_cells('A2:J2')
    ws['A2'] = f"Supplier: {supplier or '—'}    Date: {timezone.localdate().strftime('%d-%m-%Y')}"
    ws['A2'].font = Font(bold=True, size=11)
    ws['A2'].alignment = center

    if sheet.get('note'):
        ws.merge_cells('A3:J3')
        ws['A3'] = sheet['note']
        ws['A3'].alignment = center

    headers = ['Sr', 'Date', 'Product / Item', 'Qty', 'Tax %',
               'Rate (Excl Tax)', 'Rate (Incl Tax)', 'Amount (Incl Tax)',
               'Available? (Supplier fills)', 'Supplier Rate (Supplier fills)']
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=5, column=col, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = border

    row = 6
    for idx, r in enumerate(rows, 1):
        has_excl = r['price_excl'] is not None
        has_incl = r['price_incl'] is not None
        if r['price_src'] == 'incl' and has_incl:
            excl_cell = f'=ROUND(G{row}/(1+E{row}/100),2)'
            incl_cell = float(r['price_incl'])
        elif has_excl:
            excl_cell = float(r['price_excl'])
            incl_cell = f'=ROUND(F{row}*(1+E{row}/100),2)'
        elif has_incl:
            excl_cell = f'=ROUND(G{row}/(1+E{row}/100),2)'
            incl_cell = float(r['price_incl'])
        else:
            excl_cell = incl_cell = None
        amount_cell = f'=ROUND(D{row}*G{row},2)' if (has_excl or has_incl) else None

        values = [
            idx,
            r['date'],
            r['product'],
            float(r['qty']) if r['qty'] else None,
            float(r['tax']),
            excl_cell,
            incl_cell,
            amount_cell,
            None, None,
        ]
        for col, v in enumerate(values, 1):
            c = ws.cell(row=row, column=col, value=v)
            c.border = border
            if col in (4, 5, 6, 7, 8):
                c.alignment = right_a
                if col in (6, 7, 8):
                    c.number_format = money_fmt

        for col in (9, 10):
            ws.cell(row=row, column=col).fill = PatternFill(
                start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
        row += 1

    last = row - 1
    label_row = row
    grand_row = label_row + 2
    ws.cell(row=label_row, column=3, value='Subtotal (Excl Tax)').font = Font(bold=True)
    c = ws.cell(row=label_row, column=6,
                value=f'=ROUND(SUMPRODUCT(D6:D{last},F6:F{last}),2)')
    c.font = Font(bold=True); c.number_format = money_fmt; c.alignment = right_a
    ws.cell(row=label_row + 1, column=3, value='Total Tax').font = Font(bold=True)
    c = ws.cell(row=label_row + 1, column=6, value=f'=ROUND(H{grand_row}-F{label_row},2)')
    c.font = Font(bold=True); c.number_format = money_fmt; c.alignment = right_a
    ws.cell(row=grand_row, column=3, value='GRAND TOTAL').font = Font(bold=True, size=12)
    c = ws.cell(row=grand_row, column=8, value=f'=ROUND(SUM(H6:H{last}),2)')
    c.font = Font(bold=True, size=12); c.number_format = money_fmt; c.alignment = right_a
    for rr in (label_row, label_row + 1, grand_row):
        for col in range(1, 11):
            ws.cell(row=rr, column=col).border = border

    widths = [5, 12, 34, 8, 8, 15, 15, 16, 22, 22]
    for idx, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    ws.freeze_panes = 'A6'

    slug = re.sub(r'[^a-z0-9]+', '-', supplier.lower()).strip('-') or 'supplier'
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = (
        f'attachment; filename="order-note_{slug}_'
        f'{timezone.localdate().strftime("%Y%m%d")}.xlsx"'
    )
    wb.save(response)
    return response
