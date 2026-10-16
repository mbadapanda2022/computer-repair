# accounting/views/payments.py

import json
import logging
from decimal import Decimal, InvalidOperation

from django.shortcuts import render, redirect, get_object_or_404
from django.http import HttpResponse
from django.contrib import messages
from django.db import transaction
from django.db.models import Q, Sum, F
from django.core.paginator import Paginator, EmptyPage, PageNotAnInteger
from django.views.decorators.csrf import csrf_protect
from django.views.decorators.http import require_http_methods
from django.urls import reverse

from ..models import (
    Payment, BankTransaction, Contact, Invoice, PaymentAllocation,
    Purchase, PurchaseAllocation, CompanyProfile,
)
from ..forms import PaymentForm
from .utils import is_htmx, htmx_response, redirect_to_staff
from ..decorators import handle_errors

from accounting.utils.notification_helpers import (
    send_notification_to_customer,
    send_notification_to_staff,
)

logger = logging.getLogger(__name__)


# ============================================================
# HELPERS
# ============================================================
def get_unpaid_invoices_for_contact(contact, direction='received'):
    """
    Return unpaid or partially-paid invoices for a customer contact.
    Only meaningful for direction='received'.
    """
    if direction == 'received' and contact and contact.contact_type in ('customer', 'both'):
        return (
            Invoice.objects
            .filter(customer=contact, balance_due__gt=0)
            .exclude(payment_status='paid')
            .order_by('date', 'id')
        )
    return Invoice.objects.none()


def _safe_decimal(value, default=Decimal('0')):
    if value is None:
        return default
    s = str(value).strip()
    if s == '':
        return default
    try:
        return Decimal(s)
    except (ValueError, InvalidOperation, TypeError):
        return default


def _safe_next_url(request):
    """Internal-path redirect target from ?next= (blocks off-site redirects)."""
    candidate = request.POST.get('next') or request.GET.get('next')
    if candidate and candidate.startswith('/') and not candidate.startswith('//'):
        return candidate
    return None


def _filter_payment_queryset(request):
    """Apply list filters from GET params. Returns (queryset, params)."""
    queryset = (
        Payment.objects
        .select_related('contact', 'bank_account', 'expense_account')
        .prefetch_related('allocations__invoice', 'purchase_allocations__purchase')
        .all()
        .order_by('-date', '-id')
    )

    params = {
        'search': request.GET.get('search', '').strip(),
        'direction': request.GET.get('direction', '').strip(),
        'method': request.GET.get('method', '').strip(),
        'date_from': request.GET.get('date_from', '').strip(),
        'date_to': request.GET.get('date_to', '').strip(),
        'reconciled': request.GET.get('reconciled', '').strip(),
    }

    if params['search']:
        queryset = queryset.filter(
            Q(contact__name__icontains=params['search']) |
            Q(contact__phone__icontains=params['search']) |
            Q(reference__icontains=params['search']) |
            Q(upi_ref__icontains=params['search']) |
            Q(description__icontains=params['search']) |
            Q(expense_account__name__icontains=params['search']) |
            Q(allocations__invoice__invoice_number__icontains=params['search']) |
            Q(purchase_allocations__purchase__purchase_number__icontains=params['search'])
        ).distinct()
    if params['direction']:
        queryset = queryset.filter(direction=params['direction'])
    if params['method']:
        queryset = queryset.filter(method=params['method'])
    if params['date_from']:
        queryset = queryset.filter(date__gte=params['date_from'])
    if params['date_to']:
        queryset = queryset.filter(date__lte=params['date_to'])
    if params['reconciled'] in ('yes', 'no'):
        queryset = queryset.filter(reconciled=(params['reconciled'] == 'yes'))

    return queryset, params


def _get_paginated_payments_context(request, queryset=None):
    """Shared filter + paginate + annotate for payments list."""
    if queryset is None:
        queryset, params = _filter_payment_queryset(request)
    else:
        params = {
            'search': request.GET.get('search', '').strip(),
            'direction': request.GET.get('direction', '').strip(),
            'method': request.GET.get('method', '').strip(),
            'date_from': request.GET.get('date_from', '').strip(),
            'date_to': request.GET.get('date_to', '').strip(),
            'reconciled': request.GET.get('reconciled', '').strip(),
        }
    page_number = request.GET.get('page', 1)

    # ── Totals (computed on filtered queryset) ──
    # Note: `.distinct()` + `.aggregate()` is safe here because distinct
    # collapses duplicate rows from JOIN, and SUM runs over the deduped set.
    total_received = queryset.filter(direction='received').aggregate(
        t=Sum('amount')
    )['t'] or Decimal('0')
    total_paid = queryset.filter(direction='paid').aggregate(
        t=Sum('amount')
    )['t'] or Decimal('0')
    total_amount = queryset.aggregate(t=Sum('amount'))['t'] or Decimal('0')
    upi_count = queryset.filter(method='upi').count()
    pending_reconciliation = queryset.filter(reconciled=False).count()

    # ── Pagination ──
    paginator = Paginator(queryset, 15)
    try:
        page_obj = paginator.page(page_number)
    except PageNotAnInteger:
        page_obj = paginator.page(1)
    except EmptyPage:
        page_obj = paginator.page(paginator.num_pages)

    from urllib.parse import urlencode
    active_parts = {k: v for k, v in params.items() if v}
    filters_qs = ('&' + urlencode(active_parts)) if active_parts else ''

    return {
        'payments': page_obj.object_list,
        'page_obj': page_obj,
        'total_received': total_received,
        'total_paid': total_paid,
        'upi_count': upi_count,
        'pending_reconciliation': pending_reconciliation,
        'total_amount': total_amount,
        'filters_qs': filters_qs,
        **params,
    }


def _collect_allocations_from_post(request, payment, existing_invoice_ids=None):
    """
    Parse `allocation_amount_<invoice_id>` fields from POST and create
    PaymentAllocation rows. Returns list of description strings.

    `existing_invoice_ids` — set of invoice ids already allocated
                             (used to skip duplicates when pre-allocated).
    """
    made = []
    skip_ids = set(existing_invoice_ids or [])
    contact = payment.contact

    for key, value in request.POST.items():
        if not key.startswith('allocation_amount_'):
            continue
        if not value or not value.strip():
            continue

        try:
            inv_id = int(key.replace('allocation_amount_', ''))
            amt = _safe_decimal(value, default=Decimal('0'))
        except (ValueError, TypeError):
            continue

        if amt <= 0:
            continue

        try:
            inv = Invoice.objects.get(pk=inv_id, customer=contact)
        except Invoice.DoesNotExist:
            continue

        if inv.pk in skip_ids:
            continue

        alloc_amt = min(amt, inv.balance_due)
        if alloc_amt <= 0:
            continue

        PaymentAllocation.objects.create(
            payment=payment,
            invoice=inv,
            amount=alloc_amt,
        )
        made.append(f"{inv.invoice_number} Rs.{alloc_amt}")
        skip_ids.add(inv.pk)

    return made


def get_unpaid_purchases_for_contact(contact):
    """Vendor bills with a balance due for a vendor contact."""
    if contact and contact.contact_type in ('vendor', 'both'):
        return (
            Purchase.objects
            .filter(vendor=contact, paid_amount__lt=F('grand_total'))
            .order_by('date', 'id')
        )
    return Purchase.objects.none()


def _collect_purchase_allocations_from_post(request, payment, existing_purchase_ids=None):
    """
    Parse `purchase_alloc_<purchase_id>` fields from POST and create
    PurchaseAllocation rows. Returns list of description strings.
    """
    made = []
    skip_ids = set(existing_purchase_ids or [])
    contact = payment.contact

    for key, value in request.POST.items():
        if not key.startswith('purchase_alloc_'):
            continue
        if not value or not value.strip():
            continue

        try:
            p_id = int(key.replace('purchase_alloc_', ''))
            amt = _safe_decimal(value, default=Decimal('0'))
        except (ValueError, TypeError):
            continue

        if amt <= 0:
            continue

        try:
            purchase = Purchase.objects.get(pk=p_id, vendor=contact)
        except Purchase.DoesNotExist:
            continue

        if purchase.pk in skip_ids:
            continue

        alloc_amt = min(amt, purchase.balance_due)
        if alloc_amt <= 0:
            continue

        PurchaseAllocation.objects.create(
            payment=payment,
            purchase=purchase,
            amount=alloc_amt,
        )
        made.append(f"{purchase.purchase_number} Rs.{alloc_amt}")
        skip_ids.add(purchase.pk)

    return made


# ============================================================
# LIST
# ============================================================
@handle_errors(default_redirect='accounting:payment_list')
def payment_list(request):
    context = _get_paginated_payments_context(request)
    if is_htmx(request):
        return render(request, 'payments/partials/payment_table.html', context)
    return render(request, 'payments/payment_list.html', context)


# ============================================================
# PAYMENT DETAIL MODAL
# ============================================================
def payment_detail_modal(request, pk):
    payment = get_object_or_404(
        Payment.objects.select_related(
            'contact', 'bank_account', 'expense_account',
        ).prefetch_related(
            'allocations__invoice', 'advance_adjustments__invoice',
            'purchase_allocations__purchase',
        ),
        pk=pk,
    )
    allocated_total = sum(
        (a.amount for a in payment.allocations.all()), Decimal('0')
    ) + sum(
        (a.amount for a in payment.advance_adjustments.all()), Decimal('0')
    ) + sum(
        (a.amount for a in payment.purchase_allocations.all()), Decimal('0')
    )
    unallocated = payment.amount - allocated_total
    return render(request, 'payments/partials/payment_detail_modal.html', {
        'payment': payment,
        'allocated_total': allocated_total,
        'unallocated': unallocated,
    })


# ============================================================
# PAYMENT PRINT (single receipt voucher)
# ============================================================
def payment_print(request, pk):
    payment = get_object_or_404(
        Payment.objects.select_related(
            'contact', 'bank_account', 'expense_account',
        ).prefetch_related(
            'allocations__invoice', 'advance_adjustments__invoice',
            'purchase_allocations__purchase',
        ),
        pk=pk,
    )
    company = CompanyProfile.get_instance()
    logo_exists = bool(
        company.logo and company.logo.name and company.logo.storage.exists(company.logo.name)
    )
    return render(request, 'payments/payment_print.html', {
        'payment': payment,
        'company': company,
        'logo_exists': logo_exists,
    })


# ============================================================
# PAYMENT LIST — PRINT (filtered, no pagination)
# ============================================================
@handle_errors(default_redirect='accounting:payment_list')
def payment_list_print(request):
    queryset, params = _filter_payment_queryset(request)
    payments = list(queryset)
    company = CompanyProfile.get_instance()

    total_received = sum(
        (p.amount for p in payments if p.direction == 'received'), Decimal('0')
    )
    total_paid = sum(
        (p.amount for p in payments if p.direction == 'paid'), Decimal('0')
    )

    return render(request, 'payments/payment_list_print.html', {
        'payments': payments,
        'company': company,
        'total_received': total_received,
        'total_paid': total_paid,
        **params,
    })


# ============================================================
# PAYMENT LIST — EXCEL EXPORT (filtered)
# ============================================================
@handle_errors(default_redirect='accounting:payment_list')
def payment_list_excel(request):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    queryset, params = _filter_payment_queryset(request)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Payments"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    title_font = Font(bold=True, size=14, color="1F4E78")
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'),
    )
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    right = Alignment(horizontal='right', vertical='center')
    money_fmt = '#,##0.00'

    headers = [
        'Date', 'Payment ID', 'Direction', 'Contact', 'Amount (Rs.)',
        'Discount (Rs.)', 'Net (Rs.)', 'Method', 'UPI / Ref', 'Account',
        'Applied To', 'Expense Head', 'Advance', 'Reconciled', 'Description',
    ]
    TOTAL_COLS = len(headers)
    last_col = get_column_letter(TOTAL_COLS)

    ws.merge_cells(f'A1:{last_col}1')
    t = ws.cell(row=1, column=1, value="Payment Register")
    t.font = title_font
    t.alignment = center
    ws.row_dimensions[1].height = 30

    for col_idx, header in enumerate(headers, 1):
        c = ws.cell(row=2, column=col_idx, value=header)
        c.font = header_font
        c.fill = header_fill
        c.border = thin_border
        c.alignment = center
    ws.row_dimensions[2].height = 25

    for row_idx, p in enumerate(queryset, start=3):
        bill_parts = [
            f"{a.invoice.invoice_number}({a.amount})"
            for a in p.allocations.select_related('invoice').all()
        ] + [
            f"{a.purchase.purchase_number}({a.amount})"
            for a in p.purchase_allocations.select_related('purchase').all()
        ]
        applied = ', '.join(bill_parts) or ', '.join(
            f"ADV:{a.invoice.invoice_number}({a.amount})"
            for a in p.advance_adjustments.select_related('invoice').all()
        )
        row_data = [
            p.date.strftime('%d-%m-%Y'),
            p.pk,
            p.get_direction_display(),
            p.contact.name if p.contact else '',
            float(p.amount),
            float(p.discount_amount or 0),
            float(p.net_amount),
            p.get_method_display(),
            p.upi_ref or p.reference or '',
            p.bank_account.name if p.bank_account else 'Cash',
            applied,
            p.expense_account.name if p.expense_account else '',
            'Yes' if p.is_advance else '',
            'Yes' if p.reconciled else 'Pending',
            p.description or '',
        ]
        for col_idx, value in enumerate(row_data, 1):
            c = ws.cell(row=row_idx, column=col_idx, value=value)
            c.border = thin_border
            if col_idx in (5, 6, 7):
                c.alignment = right
                c.number_format = money_fmt
            elif col_idx in (1, 2, 13, 14):
                c.alignment = center
            else:
                c.alignment = left

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = max(min(max_len + 3, 45), 10)

    ws.freeze_panes = 'A3'

    from django.utils import timezone
    fname = f"payments_{timezone.localdate().strftime('%Y%m%d')}.xlsx"
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{fname}"'
    wb.save(response)
    return response


# ============================================================
# LOAD UNPAID BILLS (HTMX endpoint) — customer invoices OR vendor purchase
# bills, depending on payment direction.
# ============================================================
@require_http_methods(["GET"])
def load_unpaid_invoices(request):
    """Return partial HTML with unpaid invoices/purchases for a contact."""
    contact_id = request.GET.get('contact', '').strip()
    direction = request.GET.get('direction', 'received').strip()
    payment_pk = request.GET.get('payment', '').strip()

    if not contact_id:
        return HttpResponse('')

    try:
        contact = Contact.objects.get(pk=contact_id)
    except (Contact.DoesNotExist, ValueError, TypeError):
        return HttpResponse('')

    payment = None
    if payment_pk:
        try:
            payment = Payment.objects.get(pk=payment_pk)
        except (Payment.DoesNotExist, ValueError, TypeError):
            payment = None

    if direction == 'received':
        if payment:
            unpaid = _unpaid_invoices_for_edit(payment)
        else:
            unpaid = get_unpaid_invoices_for_contact(contact, 'received')
        return render(request, 'payments/partials/unpaid_invoices.html', {
            'contact': contact,
            'unpaid_invoices': unpaid,
            'direction': direction,
            'payment': payment,
        })

    if payment:
        unpaid_purchases = _unpaid_purchases_for_edit(payment)
    else:
        unpaid_purchases = get_unpaid_purchases_for_contact(contact)
    return render(request, 'payments/partials/unpaid_purchases.html', {
        'contact': contact,
        'unpaid_purchases': unpaid_purchases,
        'direction': direction,
        'payment': payment,
    })


# ============================================================
# CREATE
# ============================================================
@csrf_protect
@handle_errors(default_redirect='accounting:payment_list',
               htmx_template='payments/payment_form.html')
def payment_create(request, pk=None):
    """
    Create a new payment.

    Prefill support:
      - ?contact=<id>  → prefill contact, set direction from contact type
      - ?invoice=<id>  → prefill contact + amount, auto-allocate on save

    Multi-invoice allocation via POST fields:
      - allocation_amount_<invoice_id> = "123.45"
    """
    initial = {}
    contact_id = request.GET.get('contact')
    invoice_id = request.GET.get('invoice')
    purchase_id = request.GET.get('purchase')
    if pk is not None and not invoice_id:
        invoice_id = pk

    # Prefill from contact query param
    if contact_id:
        try:
            c = Contact.objects.get(pk=contact_id)
            initial['contact'] = c.id
            initial['direction'] = 'received' if c.contact_type in ('customer', 'both') else 'paid'
        except (Contact.DoesNotExist, ValueError, TypeError):
            pass

    # Prefill from invoice
    target_invoice = None
    if invoice_id:
        try:
            target_invoice = Invoice.objects.select_related('customer').get(pk=invoice_id)
            initial['contact'] = target_invoice.customer_id
            initial['direction'] = 'received'
            initial['amount'] = target_invoice.balance_due
        except (Invoice.DoesNotExist, ValueError, TypeError):
            pass

    # Prefill from purchase bill (vendor settlement)
    target_purchase = None
    if purchase_id:
        try:
            target_purchase = Purchase.objects.select_related('vendor').get(pk=purchase_id)
            initial['contact'] = target_purchase.vendor_id
            initial['direction'] = 'paid'
            initial['amount'] = target_purchase.balance_due
        except (Purchase.DoesNotExist, ValueError, TypeError):
            pass

    # ============================================================
    # POST
    # ============================================================
    if request.method == 'POST':
        form = PaymentForm(request.POST)
        if form.is_valid():
            allocation_summary = []

            # ── 1. Save payment atomically ──
            with transaction.atomic():
                payment = form.save()

                if not payment.is_advance and not payment.expense_account_id:
                    # Auto-allocate to target invoice (from ?invoice=)
                    if target_invoice and target_invoice.customer_id == payment.contact_id:
                        alloc_amt = min(payment.amount, target_invoice.balance_due)
                        if alloc_amt > 0:
                            PaymentAllocation.objects.create(
                                payment=payment,
                                invoice=target_invoice,
                                amount=alloc_amt,
                            )
                            allocation_summary.append(
                                f"{target_invoice.invoice_number} Rs.{alloc_amt}"
                            )

                    # Auto-allocate to target purchase bill (from ?purchase=)
                    if target_purchase and target_purchase.vendor_id == payment.contact_id:
                        alloc_amt = min(payment.net_amount, target_purchase.balance_due)
                        if alloc_amt > 0:
                            PurchaseAllocation.objects.create(
                                payment=payment,
                                purchase=target_purchase,
                                amount=alloc_amt,
                            )
                            allocation_summary.append(
                                f"{target_purchase.purchase_number} Rs.{alloc_amt}"
                            )

                    # Multi-invoice allocations from POST
                    allocation_summary += _collect_allocations_from_post(
                        request, payment,
                        existing_invoice_ids=[target_invoice.pk] if target_invoice else [],
                    )

                    # Vendor bill allocations from POST
                    allocation_summary += _collect_purchase_allocations_from_post(
                        request, payment,
                        existing_purchase_ids=(
                            [target_purchase.pk] if target_purchase else []
                        ),
                    )

                    if allocation_summary:
                        payment.update_invoices()
                        payment.update_purchases()

            logger.info(
                "Payment #%s created by %s (Rs.%s %s). Allocations: %s",
                payment.pk, request.user.username, payment.amount,
                payment.direction, allocation_summary or 'None',
            )

            # ── 2. Notifications OUTSIDE atomic block ──
            try:
                if payment.direction == 'received' and payment.contact:
                    send_notification_to_customer(
                        payment.contact,
                        title=f"Payment Received: Rs.{payment.amount}",
                        message=f"Your payment of Rs.{payment.amount} has been recorded.",
                        link=reverse('customer:customer_payments'),
                        notif_type='success',
                        category='payment',
                        send_email=False,
                    )

                send_notification_to_staff(
                    title=f"New Payment #{payment.pk}",
                    message=(
                        f"{payment.get_direction_display()} of "
                        f"Rs.{payment.amount} with {payment.contact.name}"
                    ),
                    link=reverse('accounting:payment_list'),
                    notif_type='info',
                    category='payment',
                    send_email=False,
                )
            except Exception:
                logger.exception("Payment notification failed for #%s", payment.pk)

            # ── 3. Response ──
            if is_htmx(request):
                context = _get_paginated_payments_context(request)
                response = render(request, 'payments/partials/payment_table.html', context)
                response['HX-Trigger'] = json.dumps({
                    'showToast': {
                        'level': 'success',
                        'message': f'Payment Rs.{payment.amount} recorded successfully.',
                    },
                    'closeModal': '',
                })
                return response

            messages.success(request, f"Payment Rs.{payment.amount} recorded successfully.")
            next_url = _safe_next_url(request)
            if next_url:
                return redirect(next_url)
            return redirect_to_staff('payment_list')

        # ── Invalid form — re-render with full context ──
        if is_htmx(request):
            pre_contact = None
            pre_unpaid = Invoice.objects.none()
            pre_unpaid_purchases = Purchase.objects.none()
            try:
                pre_contact = form.data.get('contact') and Contact.objects.get(pk=form.data.get('contact'))
                if pre_contact and form.data.get('direction') == 'received':
                    pre_unpaid = get_unpaid_invoices_for_contact(pre_contact, 'received')
                elif pre_contact and form.data.get('direction') == 'paid':
                    pre_unpaid_purchases = get_unpaid_purchases_for_contact(pre_contact)
            except (Contact.DoesNotExist, ValueError, TypeError):
                pass

            response = render(request, 'payments/payment_form.html', {
                'form': form,
                'payment': None,
                'target_invoice': target_invoice,
                'target_purchase': target_purchase,
                'pre_contact': pre_contact,
                'unpaid_invoices': pre_unpaid,
                'unpaid_purchases': pre_unpaid_purchases,
            })
            response['HX-Retarget'] = '#mainModalContent'
            return response

        # Non-HTMX (native submit from pages without #payment-table-container):
        # bounce back to the caller instead of rendering a bare partial page.
        first_error = next(iter(e[0] for e in form.errors.values() if e), None)
        messages.error(request, first_error or "Please correct the payment details.")
        next_url = _safe_next_url(request)
        if next_url:
            return redirect(next_url)
        return redirect_to_staff('payment_list')

    else:
        form = PaymentForm(initial=initial)

    # ============================================================
    # GET render
    # ============================================================
    unpaid_invoices = Invoice.objects.none()
    unpaid_purchases = Purchase.objects.none()
    pre_contact = None

    initial_contact_id = form.initial.get('contact')
    if initial_contact_id:
        try:
            pre_contact = Contact.objects.get(pk=initial_contact_id)
            if form.initial.get('direction') == 'received':
                unpaid_invoices = get_unpaid_invoices_for_contact(pre_contact, 'received')
            elif form.initial.get('direction') == 'paid':
                unpaid_purchases = get_unpaid_purchases_for_contact(pre_contact)
        except (Contact.DoesNotExist, ValueError, TypeError):
            pass

    return render(request, 'payments/payment_form.html', {
        'form': form,
        'payment': None,
        'target_invoice': target_invoice,
        'target_purchase': target_purchase,
        'unpaid_invoices': unpaid_invoices,
        'unpaid_purchases': unpaid_purchases,
        'pre_contact': pre_contact,
    })


# ============================================================
# UPDATE
# ============================================================
@csrf_protect
@handle_errors(default_redirect='accounting:payment_list',
               htmx_template='payments/payment_form.html')
def payment_update(request, pk):
    """Update an existing payment (with optional allocation reset)."""
    payment = get_object_or_404(Payment, pk=pk)

    if request.method == 'POST':
        form = PaymentForm(request.POST, instance=payment)
        if form.is_valid():
            with transaction.atomic():
                payment = form.save()

                # Reset allocations ONLY if flag was explicitly set
                # (form JS sets it to '1' on submit)
                reset_flag = request.POST.get('reset_allocations', '0') == '1'

                # A direct-expense payment never settles invoices/bills; if it
                # was converted to one, drop old allocations so paid_amounts
                # don't keep crediting a payment that no longer lives there.
                if payment.expense_account_id and (
                    payment.allocations.exists() or payment.purchase_allocations.exists()
                ):
                    for alloc in list(payment.allocations.all()):
                        alloc.hard_delete()
                    for alloc in list(payment.purchase_allocations.all()):
                        alloc.hard_delete()
                    payment.update_invoices()
                    payment.update_purchases()
                elif payment.direction == 'received' and payment.purchase_allocations.exists():
                    # A customer receipt cannot settle vendor bills
                    for alloc in list(payment.purchase_allocations.all()):
                        alloc.hard_delete()
                    payment.update_purchases()
                elif payment.direction == 'paid' and payment.allocations.exists():
                    # A vendor payment cannot settle sales invoices
                    for alloc in list(payment.allocations.all()):
                        alloc.hard_delete()
                    payment.update_invoices()
                elif not payment.is_advance and not payment.expense_account_id and reset_flag:
                    # HARD delete — unique constraint (payment, invoice/purchase)
                    # does NOT respect soft-delete; soft rows would block
                    # re-creation.
                    for alloc in list(payment.allocations.all()):
                        alloc.hard_delete()
                    for alloc in list(payment.purchase_allocations.all()):
                        alloc.hard_delete()

                    _collect_allocations_from_post(request, payment)
                    _collect_purchase_allocations_from_post(request, payment)
                    payment.update_invoices()
                    payment.update_purchases()

            logger.info("Payment #%s updated by %s", pk, request.user.username)

            # Notifications OUTSIDE atomic
            try:
                if payment.direction == 'received' and payment.contact:
                    send_notification_to_customer(
                        payment.contact,
                        title=f"Payment Updated: Rs.{payment.amount}",
                        message=f"Your payment of Rs.{payment.amount} has been updated.",
                        link=reverse('customer:customer_payments'),
                        notif_type='info',
                        category='payment',
                        send_email=False,
                    )
            except Exception:
                logger.exception("Payment update notification failed for #%s", pk)

            if is_htmx(request):
                context = _get_paginated_payments_context(request)
                response = render(request, 'payments/partials/payment_table.html', context)
                response['HX-Trigger'] = json.dumps({
                    'showToast': {'level': 'success', 'message': 'Payment updated successfully.'},
                    'closeModal': '',
                })
                return response

            messages.success(request, "Payment updated successfully.")
            return redirect_to_staff('payment_list')

        # Invalid form — re-render with full context
        if is_htmx(request):
            unpaid_invoices = _unpaid_invoices_for_edit(payment)
            unpaid_purchases = _unpaid_purchases_for_edit(payment)
            response = render(request, 'payments/payment_form.html', {
                'form': form,
                'payment': payment,
                'unpaid_invoices': unpaid_invoices,
                'unpaid_purchases': unpaid_purchases,
                'pre_contact': payment.contact,
            })
            response['HX-Retarget'] = '#mainModalContent'
            return response

    else:
        form = PaymentForm(instance=payment)

    # GET: render edit form
    unpaid_invoices = _unpaid_invoices_for_edit(payment)
    unpaid_purchases = _unpaid_purchases_for_edit(payment)

    return render(request, 'payments/payment_form.html', {
        'form': form,
        'payment': payment,
        'target_invoice': None,
        'unpaid_invoices': unpaid_invoices,
        'unpaid_purchases': unpaid_purchases,
        'pre_contact': payment.contact,
    })


def _unpaid_purchases_for_edit(payment):
    """
    Return vendor bills for the edit form:
      - bills with a balance due for the vendor, PLUS
      - already-allocated bills (even if fully settled)
    """
    if not payment.contact or payment.direction != 'paid':
        return Purchase.objects.none()
    if payment.is_advance or payment.expense_account_id:
        return Purchase.objects.none()

    allocated_ids = list(payment.purchase_allocations.values_list('purchase_id', flat=True))

    return (
        Purchase.objects
        .filter(vendor=payment.contact)
        .filter(Q(paid_amount__lt=F('grand_total')) | Q(pk__in=allocated_ids))
        .order_by('date', 'id')
    )


def _unpaid_invoices_for_edit(payment):
    """
    Return invoices for the edit form:
      - unpaid/partial invoices for the customer, PLUS
      - already-allocated invoices (even if fully paid)
    """
    if not payment.contact or payment.direction != 'received':
        return Invoice.objects.none()

    allocated_ids = list(payment.allocations.values_list('invoice_id', flat=True))

    return (
        Invoice.objects
        .filter(customer=payment.contact)
        .filter(Q(balance_due__gt=0) | Q(pk__in=allocated_ids))
        .order_by('date', 'id')
    )


# ============================================================
# DELETE
# ============================================================
@csrf_protect
@require_http_methods(["DELETE"])
@handle_errors(default_redirect='accounting:payment_list')
def payment_delete(request, pk):
    payment = get_object_or_404(Payment, pk=pk)
    contact = payment.contact
    amount = payment.amount
    direction = payment.direction

    payment.delete()  # model-level cleanup
    logger.info("Payment #%s deleted by %s", pk, request.user.username)

    # Notify staff AFTER delete
    try:
        send_notification_to_staff(
            title=f"Payment Deleted: Rs.{amount}",
            message=(
                f"{direction} payment of Rs.{amount} from "
                f"{contact.name if contact else 'Unknown'} was deleted."
            ),
            link=reverse('accounting:payment_list'),
            notif_type='warning',
            category='payment',
            send_email=False,
        )
    except Exception:
        logger.exception("Payment deletion notification failed for #%s", pk)

    if is_htmx(request):
        context = _get_paginated_payments_context(request)
        response = render(request, 'payments/partials/payment_table.html', context)
        response['HX-Trigger'] = json.dumps({
            'showToast': {'level': 'success', 'message': 'Payment deleted successfully.'}
        })
        return response

    messages.success(request, "Payment deleted.")
    return redirect_to_staff('payment_list')


# ============================================================
# RECONCILE TOGGLE — LIGHTWEIGHT
# ============================================================
@csrf_protect
@require_http_methods(["POST"])
@handle_errors(default_redirect='accounting:payment_list')
def reconcile_payment(request, pk):
    """
    Toggle reconciliation status — LIGHTWEIGHT.

    Uses queryset.update() + manual bank txn sync to avoid rebuilding
    ledger entries / allocations / contact balances. Previously this
    called payment.save() which rebuilt everything on every toggle.
    """
    payment = get_object_or_404(Payment, pk=pk)

    # Accept explicit target state if provided (avoids double-toggle races)
    target = request.POST.get('reconciled')
    if target in ('true', '1', 'on'):
        new_value = True
    elif target in ('false', '0', 'off'):
        new_value = False
    else:
        new_value = not payment.reconciled

    # Fast DB update — no signals, no ledger rebuild
    Payment.objects.filter(pk=pk).update(reconciled=new_value)

    # Mirror to bank transaction (lightweight single-row update)
    BankTransaction.objects.filter(payment_id=pk).update(reconciled=new_value)

    logger.info(
        "Payment #%s reconciliation toggled to %s by %s",
        pk, new_value, request.user.username,
    )

    if is_htmx(request):
        context = _get_paginated_payments_context(request)
        response = render(request, 'payments/partials/payment_table.html', context)
        response['HX-Trigger'] = json.dumps({
            'showToast': {
                'level': 'info',
                'message': f"Payment #{payment.pk} reconciliation updated.",
            }
        })
        return response

    messages.success(request, "Reconciliation status updated.")
    return redirect_to_staff('payment_list')